"""Streaming import of hackathon KIP, LIMS and PAK source files."""

from __future__ import annotations

import hashlib
import json
import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

SCHEMA = pa.schema(
    [
        ("metric_id", pa.string()),
        ("timestamp", pa.timestamp("ns")),
        ("value", pa.float64()),
        ("source", pa.string()),
        ("unit", pa.string()),
        ("flags", pa.string()),
        ("source_file", pa.string()),
        ("source_row", pa.int64()),
    ]
)
PRIMARY = {"Mg.Sulfur", "D15", "FlashPoint", "CFPP", "CloudPoint", "95%.T"}


def _text(value: Any) -> str | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    value = str(value).strip()
    return value or None


def _time(value: Any) -> tuple[pd.Timestamp | None, bool]:
    try:
        result = pd.to_datetime(value, errors="coerce")
        if pd.isna(result):
            return None, False
        result = pd.Timestamp(result)
        return (None, True) if result.tzinfo is not None else (result, False)
    except (TypeError, ValueError):
        return None, False


def _number(value: Any) -> tuple[float | None, bool]:
    try:
        result = float(str(value).strip().replace(",", "."))
        return (result, False) if math.isfinite(result) else (None, True)
    except (TypeError, ValueError):
        return None, True


def _plant(text: str) -> str:
    low = text.lower()
    return "ht" if "гидро" in low or "24-2000" in low else "avt"


def _registry_data() -> dict[str, Any]:
    path = Path(__file__).parent / "resources" / "registry.json"
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {"metrics": [], "formulas": []}
    return raw


class Sink:
    def __init__(self, path: Path):
        self.writer = pq.ParquetWriter(path, SCHEMA, compression="zstd")
        data = _registry_data()
        self.registry = {str(x.get("id")): x for x in data.get("metrics", []) if x.get("id")}
        self.formulas = data.get("formulas", [])
        self.identities: set[str] = set()
        self.metrics: dict[str, dict[str, Any]] = {}
        self.sources: list[dict[str, Any]] = []
        self.issues: dict[tuple[str, str | None], int] = {}
        self.start: pd.Timestamp | None = None
        self.end: pd.Timestamp | None = None
        self.telemetry_start: pd.Timestamp | None = None
        self.telemetry_end: pd.Timestamp | None = None

    def metric(
        self,
        metric_id: str,
        source: str,
        unit: str | None,
        plant: str,
        code: str,
        group: str,
        description: str | None = None,
        point: str | None = None,
    ):
        if metric_id in self.metrics:
            return
        base = {
            "id": metric_id,
            "label": description or code,
            "source": source,
            "unit": unit,
            "unit_status": "confirmed" if unit else "unknown",
            "plant": plant,
            "code": code,
            "group": group,
        }
        if unit:
            base["source_unit"] = unit
        if description:
            base["description"] = description
        if point:
            base["point"] = point
        if source == "lims" and plant == "ht" and point == "2" and code in PRIMARY:
            base["primary"] = True
        override = self.registry.get(metric_id, {})
        for key in (
            "label",
            "description",
            "unit",
            "unit_status",
            "primary",
            "mapping_warning",
            "group",
        ):
            if description and key in {"label", "description"}:
                continue
            if key in override:
                base[key] = override[key]
        self.metrics[metric_id] = base

    def issue(self, code: str, metric_id: str | None = None, count: int = 1):
        self.issues[(code, metric_id)] = self.issues.get((code, metric_id), 0) + count

    def write(self, records: list[dict[str, Any]], source: dict[str, Any]):
        if not records:
            return
        records.sort(key=lambda r: (r["timestamp"], r["source_row"]))
        output, run_start, previous_ts, previous_value = [], None, None, None
        groups: dict[pd.Timestamp, list[dict[str, Any]]] = {}
        for rec in records:
            groups.setdefault(rec["timestamp"], []).append(rec)
        for ts in sorted(groups):
            group = groups[ts]
            distinct = {(r["value"], "invalid" in r.get("_flags", [])) for r in group}
            conflict = len(distinct) > 1
            if conflict:
                self.issue("conflict", group[0]["metric_id"], len(group))
            seen = set()
            for rec in group:
                key = (rec["value"], tuple(rec.get("_flags", [])))
                if key in seen:
                    source["duplicate_count"] += 1
                    self.issue("duplicate", rec["metric_id"])
                    continue
                seen.add(key)
                flags = set(filter(None, rec.pop("_flags", [])))
                if conflict:
                    flags.add("conflict")
                value = rec["value"]
                if "invalid" in flags:
                    run_start = previous_ts = previous_value = None
                elif rec["source"] in {"kip", "pak"} and value is not None:
                    continuous = (
                        previous_ts is not None and 0 <= (ts - previous_ts).total_seconds() <= 1_200
                    )
                    if continuous and value == previous_value:
                        if run_start is not None and (ts - run_start).total_seconds() >= 21_600:
                            flags.add("flatline")
                    else:
                        run_start = ts
                    previous_ts, previous_value = ts, value
                if _suspect(rec["metric_id"], value):
                    flags.add("suspect")
                    self.issue("suspect", rec["metric_id"])
                if "flatline" in flags:
                    self.issue("flatline", rec["metric_id"])
                if flags & {"flatline", "suspect"}:
                    source["suspect_count"] += 1
                rec["flags"] = "|".join(sorted(flags))
                output.append(rec)
                source["invalid_count"] += int("invalid" in flags)
                source["rows"] += 1
                source["start"] = min(source["start"], ts) if source["start"] is not None else ts
                source["end"] = max(source["end"], ts) if source["end"] is not None else ts
                self.start = min(self.start, ts) if self.start is not None else ts
                self.end = max(self.end, ts) if self.end is not None else ts
                if rec["source"] == "kip":
                    self.telemetry_start = (
                        min(self.telemetry_start, ts) if self.telemetry_start is not None else ts
                    )
                    self.telemetry_end = (
                        max(self.telemetry_end, ts) if self.telemetry_end is not None else ts
                    )
        if output:
            self.writer.write_table(pa.Table.from_pylist(output, schema=SCHEMA))


def _suspect(metric_id: str, value: float | None) -> bool:
    if value is None:
        return False
    code = metric_id.rsplit(".", 1)[-1]
    if code in {"I250", "I350"}:
        return not 0 <= value <= 100
    if code == "D15":
        return not 500 <= value <= 1200
    if metric_id.startswith("lims.") and metric_id.endswith(".T"):
        return value > 1000
    if "Sulfur" in code:
        return value < 0
    return False


def _source(kind: str, path: Path) -> dict[str, Any]:
    return {
        "kind": kind,
        "filename": path.name,
        "rows": 0,
        "metrics": 0,
        "frame_count": 0,
        "start": None,
        "end": None,
        "invalid_count": 0,
        "suspect_count": 0,
        "duplicate_count": 0,
    }


def _load_tags(files: Iterable[Path]) -> tuple[dict[tuple[str, str], str], dict[str, str]]:
    tags: dict[tuple[str, str], str] = {}
    formulas: dict[str, str] = {}
    for path in files:
        if (
            path.suffix.lower() == ".xlsx"
            and "тег" in path.name.lower()
            and not path.name.startswith("~$")
        ):
            try:
                ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["КИП"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    for plant, d, t in (("avt", row[0], row[1]), ("ht", row[2], row[3])):
                        if _text(t):
                            tags[(plant, _text(t))] = _text(d) or _text(t)
                workbook = ws.parent
                if "ВАК" in workbook.sheetnames:
                    for row in workbook["ВАК"].iter_rows(min_row=2, values_only=True):
                        for col in range(0, len(row) - 1, 2):
                            fid, expression = _text(row[col]), _text(row[col + 1])
                            if fid and expression:
                                formulas[fid] = expression
            except Exception:
                pass
    return tags, formulas


def _csv(path: Path, sink: Sink, tags: dict[tuple[str, str], str]):
    header = pd.read_csv(path, nrows=0).columns.tolist()
    date_col = next(
        (c for c in header if str(c).strip().lower() in {"date", "timestamp", "datetime"}), None
    )
    if not date_col:
        raise ValueError("KIP CSV has no date column")
    plant = "ht" if "242000" in path.stem.lower() or "24-2000" in path.stem.lower() else "avt"
    identity = f"kip:{plant}"
    if identity in sink.identities:
        raise ValueError(f"Повторный источник КИП для установки {plant}")
    sink.identities.add(identity)
    codes = [str(c) for c in header if c != date_col and not str(c).lower().startswith("unnamed")]
    src = _source("kip", path)
    for code in codes:
        description = tags.get((plant, code))
        sink.metric(f"{plant}.{code}", "kip", None, plant, code, "КИП", description)
        if not description and f"{plant}.{code}" not in sink.registry:
            sink.metrics[f"{plant}.{code}"]["mapping_warning"] = (
                "Description not found in tags or registry"
            )
            sink.issue("unknown_mapping", f"{plant}.{code}")
    src["metrics"] = len(codes)
    frame = pd.read_csv(path, low_memory=False)
    parsed = pd.to_datetime(frame[date_col], errors="coerce")
    timezone = getattr(parsed.dt, "tz", None)
    if timezone is not None:
        sink.issue("timezone_timestamp", count=len(frame) * len(codes))
        parsed[:] = pd.NaT
    src["frame_count"] = int(parsed.nunique())
    for code in codes:
        records = []
        unit = sink.metrics[f"{plant}.{code}"].get("unit")
        for idx, (raw, ts) in enumerate(zip(frame[code], parsed), 2):
            if pd.isna(ts):
                sink.issue("invalid_timestamp", f"{plant}.{code}")
                continue
            value, invalid = _number(raw)
            records.append(
                {
                    "metric_id": f"{plant}.{code}",
                    "timestamp": ts,
                    "value": value,
                    "source": "kip",
                    "unit": unit,
                    "source_file": path.name,
                    "source_row": idx,
                    "_flags": ["invalid"] if invalid else [],
                }
            )
        sink.write(records, src)
    sink.sources.append(src)


def _lims(path: Path, sink: Sink):
    if "lims" in sink.identities:
        raise ValueError("Повторный источник ЛИМС")
    sink.identities.add("lims")
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    group_row, code_row, unit_row, _counts = [next(rows, ()) for _ in range(4)]
    pairs, current = [], None
    for col in range(0, max(len(group_row), len(code_row)), 2):
        if _text(group_row[col] if col < len(group_row) else None):
            current = _text(group_row[col])
        code = _text(code_row[col] if col < len(code_row) else None)
        if not current or not code:
            continue
        plant = _plant(current)
        match = re.search(r"Точка отбора\s*'([^']+)'", current, re.I)
        point = match.group(1) if match else "unknown"
        unit = _text(unit_row[col] if col < len(unit_row) else None)
        mid = f"lims.{plant}.{point}.{code}"
        pairs.append((col, mid, plant, point, code, unit, current))
        sink.metric(mid, "lims", unit, plant, code, current, point=point)
    src = _source("lims", path)
    src["metrics"] = len(pairs)
    series = {mid: [] for _, mid, *_ in pairs}
    frames = set()
    for row_no, row in enumerate(rows, 5):
        for col, mid, plant, point, code, unit, group in pairs:
            tv = row[col] if col < len(row) else None
            vv = row[col + 1] if col + 1 < len(row) else None
            if tv is None and vv is None:
                continue
            ts, zoned = _time(tv)
            if zoned:
                sink.issue("timezone_timestamp", mid)
                continue
            if ts is None:
                sink.issue("invalid_timestamp", mid)
                continue
            value, invalid = _number(vv)
            normalized_unit = sink.metrics[mid].get("unit")
            series[mid].append(
                {
                    "metric_id": mid,
                    "timestamp": ts,
                    "value": value,
                    "source": "lims",
                    "unit": normalized_unit,
                    "source_file": path.name,
                    "source_row": row_no,
                    "_flags": ["invalid"] if invalid else [],
                }
            )
            frames.add(ts)
    for records in series.values():
        sink.write(records, src)
    src["frame_count"] = len(frames)
    sink.sources.append(src)


def _pak(path: Path, sink: Sink):
    if "pak" in sink.identities:
        raise ValueError("Повторный источник ПАК")
    sink.identities.add("pak")
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    heads = next(rows, ())
    units = next(rows, ())
    pairs = []
    for col in range(0, len(heads), 1):
        head = _text(heads[col])
        if not head or col + 1 >= len(heads):
            continue
        if col and _text(heads[col - 1]):
            continue
        code = head.split(":")[-1]
        unit = _text(units[col])
        mid = f"pak.ht.{code}"
        pairs.append((col, mid, code, unit))
        sink.metric(mid, "pak", unit, "ht", code, "ПАК")
    src = _source("pak", path)
    src["metrics"] = len(pairs)
    series = {mid: [] for _, mid, *_ in pairs}
    frames = set()
    for row_no, row in enumerate(rows, 3):
        for col, mid, code, unit in pairs:
            tv = row[col] if col < len(row) else None
            vv = row[col + 1] if col + 1 < len(row) else None
            if tv is None and vv is None:
                continue
            ts, zoned = _time(tv)
            if zoned:
                sink.issue("timezone_timestamp", mid)
                continue
            if ts is None:
                sink.issue("invalid_timestamp", mid)
                continue
            value, invalid = _number(vv)
            normalized_unit = sink.metrics[mid].get("unit")
            series[mid].append(
                {
                    "metric_id": mid,
                    "timestamp": ts,
                    "value": value,
                    "source": "pak",
                    "unit": normalized_unit,
                    "source_file": path.name,
                    "source_row": row_no,
                    "_flags": ["invalid"] if invalid else [],
                }
            )
            frames.add(ts)
    for records in series.values():
        sink.write(records, src)
    src["frame_count"] = len(frames)
    sink.sources.append(src)


def _iso(value: pd.Timestamp | None) -> str | None:
    return value.isoformat() if value is not None else None


def import_dataset(files: list[Path], output_dir: Path, name: str) -> dict:
    """Import available recognized sources; source ``rows`` counts written observations."""
    output_dir.mkdir(parents=True, exist_ok=True)
    sink = Sink(output_dir / "observations.parquet")
    tags, formula_overrides = _load_tags(files)
    for formula in sink.formulas:
        fid = str(formula.get("id", ""))
        if fid in formula_overrides and formula_overrides[fid] != formula.get("expression"):
            formula["registry_expression"] = formula.get("expression")
            formula["expression"] = formula_overrides[fid]
            digest = hashlib.sha256(formula_overrides[fid].encode("utf-8")).hexdigest()[:12]
            formula["version"] = f"source-vak-{digest}"
            formula["status"] = "unresolved"
            formula["reason"] = (
                "Формула заменена значением из загруженного листа ВАК; требуется проверка"
            )
    for path in files:
        if not path.is_file() or path.name.startswith("~$"):
            continue
        try:
            if path.suffix.lower() == ".csv":
                _csv(path, sink, tags)
            elif path.suffix.lower() == ".xlsx" and "тег" not in path.name.lower():
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                first = _text(wb.active.cell(1, 1).value) or ""
                wb.close()
                if "пак" in path.name.lower() or first.startswith("24-2000:"):
                    _pak(path, sink)
                elif "лимс" in path.name.lower() or "Точка отбора" in first:
                    _lims(path, sink)
        except Exception as exc:
            sink.issue("source_error")
            sink.sources.append({**_source("unknown", path), "error": str(exc)})
    sink.writer.close()
    for source in sink.sources:
        source["start"], source["end"] = _iso(source["start"]), _iso(source["end"])
    issues = []
    messages = {
        "duplicate": "Одинаковые наблюдения метрики в один момент объединены",
        "conflict": "Конфликтующие наблюдения метрики в один момент сохранены и помечены",
        "invalid_timestamp": "Строки с некорректным временем пропущены",
        "timezone_timestamp": "Время с часовым поясом отклонено без преобразования",
        "source_error": "Источник не удалось импортировать",
        "unknown_mapping": "Описание метрики не найдено в тегах или реестре",
        "suspect": "Значение выходит за физически обоснованный диапазон контроля качества",
        "flatline": "Неизменное значение сохраняется не менее шести часов непрерывных измерений",
    }
    for (code, mid), count in sorted(sink.issues.items()):
        item = {"code": code, "message": messages.get(code, code), "count": count}
        if mid:
            item["metric_id"] = mid
        issues.append(item)
    good_rows = sum(s["rows"] for s in sink.sources)
    manifest = {
        "id": output_dir.name,
        "name": name,
        "status": "ready" if good_rows else "error",
        "created_at": datetime.now().isoformat(),
        "start": _iso(sink.start),
        "end": _iso(sink.end),
        "telemetry_start": _iso(sink.telemetry_start),
        "telemetry_end": _iso(sink.telemetry_end),
        "sources": sink.sources,
        "metrics": list(sink.metrics.values()),
        "formulas": sink.formulas,
        "issues": issues,
        "assumptions": [
            "Время источников хранится без часового пояса и без преобразования.",
            "Поле rows источника считает записанные наблюдения после объединения одинаковых дублей.",
            "Поле frame_count считает различные временные метки источника.",
            "Единицы КИП берутся только из реестра; неизвестные единицы не создаются.",
        ],
    }
    if not good_rows:
        manifest["error"] = "Не найдено пригодных для импорта наблюдений"
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest
