"""Reproducible first-pass profiling; source data stays read-only."""
from pathlib import Path
import json
import sys
import argparse
import re
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--source", required=True, type=Path)
parser.add_argument("--output", required=True, type=Path)
args = parser.parse_args()
SOURCE = args.source.resolve()
OUT = args.output.resolve()
OUT.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding="utf-8")

def dump(name, obj):
    (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

def numeric_stats(s):
    n = pd.to_numeric(s, errors="coerce").replace([np.inf, -np.inf], np.nan)
    q = n.quantile([0, .01, .05, .5, .95, .99, 1])
    runs = n.groupby(n.ne(n.shift()).cumsum()).size()
    return {"count": int(n.notna().sum()), "missing": int(n.isna().sum()), "nunique": int(n.nunique()),
            "mean": n.mean(), "std": n.std(), "min": q.iloc[0], "p01": q.iloc[1], "p05": q.iloc[2],
            "median": q.iloc[3], "p95": q.iloc[4], "p99": q.iloc[5], "max": q.iloc[6],
            "zero_count": int(n.eq(0).sum()), "negative_count": int(n.lt(0).sum()),
            "unchanged_adjacent_fraction": float(n.eq(n.shift()).mean()), "longest_constant_rows": int(runs.max()) if len(runs) else 0,
            "nonnumeric_examples": s[s.notna() & n.isna()].astype(str).value_counts().head(8).to_dict()}

def time_stats(t):
    good = t.dropna().sort_values()
    deltas = good.diff().dt.total_seconds().dropna()
    return {"start": good.min(), "end": good.max(), "invalid_dates": int(t.isna().sum()),
            "duplicate_timestamps": int(good.duplicated().sum()), "monotonic": bool(t.is_monotonic_increasing),
            "median_step_seconds": deltas.median(), "max_gap_hours": deltas.max()/3600,
            "step_counts": deltas.value_counts().head(8).to_dict()}

book = openpyxl.load_workbook(next(SOURCE.glob("Теги*")), read_only=True, data_only=True)
mapping = {}
dictionary = []
for rownum, row in enumerate(book["КИП"].iter_rows(min_row=2, values_only=True), 2):
    for stage, offset in [("avt", 0), ("242000", 2)]:
        description, tag = row[offset:offset+2]
        if tag:
            mapping[(stage, tag)] = description
            dictionary.append({"stage": stage, "tag": tag, "description_source": description, "source_cell": f"КИП!{get_column_letter(offset+2)}{rownum}", "unit_confirmed": None, "control_confirmed": False})
dump("tag_dictionary.json", dictionary)
dump("vak_source.json", [list(r) for r in book["ВАК"].values])
book.close()

telemetry = {}
profiles = []
summary = {"telemetry": {}, "pak": {}, "lims": []}
monthly = []
for path in sorted((SOURCE / "data").glob("*.csv")):
    stage = path.stem.replace("_tags", "")
    frame = pd.read_csv(path)
    indexes = [c for c in frame if c.startswith("Unnamed:")]
    time = pd.to_datetime(frame.pop("date"), errors="coerce", format="mixed")
    frame = frame.drop(columns=indexes)
    frame.index = time
    telemetry[stage] = frame
    summary["telemetry"][stage] = {"rows": len(frame), "signals": len(frame.columns), "service_columns": indexes, **time_stats(time),
        "missing_cells": int(frame.isna().sum().sum()), "missing_dictionary": [c for c in frame if (stage, c) not in mapping]}
    for col in frame:
        profiles.append({"stage": stage, "tag": col, "description": mapping.get((stage,col)), **numeric_stats(frame[col])})
    for month, group in frame.groupby(frame.index.to_period("M")):
        monthly.append({"source": stage, "month": str(month), "rows": len(group), "missing_fraction": float(group.isna().mean().mean())})
    print(stage, summary["telemetry"][stage], flush=True)
summary["telemetry_same_timestamps"] = telemetry["avt"].index.equals(telemetry["242000"].index)
pd.DataFrame(profiles).to_csv(OUT / "telemetry_profile.csv", index=False, encoding="utf-8-sig")
dump("telemetry_profile.json", profiles)

pak_raw = pd.read_excel(next(SOURCE.glob("Выгрузка*")), header=None)
pak = {}
for col, name in [(0,"sulfur"),(3,"density")]:
    frame = pd.DataFrame({"time": pd.to_datetime(pak_raw.iloc[2:, col], errors="coerce", format="mixed"), "raw": pak_raw.iloc[2:,col+1]}).dropna(subset=["time"])
    frame["value"] = pd.to_numeric(frame["raw"], errors="coerce")
    pak[name] = frame
    summary["pak"][name] = {"tag": pak_raw.iloc[0,col], "unit":pak_raw.iloc[1,col], "rows": len(frame), **time_stats(frame["time"]), **numeric_stats(frame["raw"])}
    if name == "sulfur":
        summary["pak"][name]["above_10"] = int(frame["value"].gt(10).sum())
    for month, group in frame.groupby(frame["time"].dt.to_period("M")):
        monthly.append({"source": "pak_"+name, "month": str(month), "rows": len(group), "valid": int(group["value"].notna().sum()), "median": group["value"].median(), "p95": group["value"].quantile(.95), "above_10": int(group["value"].gt(10).sum()) if name == "sulfur" else None})
    print("PAK", name, summary["pak"][name], flush=True)

labraw = pd.read_excel(next(SOURCE.glob("ЛИМС*")), header=None)
point = None
labs = []
for col in range(0, len(labraw.columns), 2):
    if pd.notna(labraw.iloc[0,col]):
        point = labraw.iloc[0,col]
    parameter = labraw.iloc[1,col]
    frame = pd.DataFrame({"time": pd.to_datetime(labraw.iloc[4:,col], errors="coerce", format="mixed"), "raw": labraw.iloc[4:,col+1]}).dropna(subset=["time"])
    frame["value"] = pd.to_numeric(frame["raw"], errors="coerce")
    series = {"point":point, "parameter":parameter, "unit_source":labraw.iloc[2,col],
              "excel_columns":f"{get_column_letter(col+1)}:{get_column_letter(col+2)}", "declared_count": int(labraw.iloc[3,col+1]),
              "rows":len(frame), **time_stats(frame["time"]), **numeric_stats(frame["raw"])}
    series["within_telemetry_valid"] = int((frame["time"].between(telemetry["avt"].index.min(),telemetry["avt"].index.max()) & frame["value"].notna()).sum())
    if parameter == "Mg.Sulfur":
        series["above_10"] = int(frame["value"].gt(10).sum())
    summary["lims"].append(series)
    labs.append((series, frame))
    for month, group in frame.groupby(frame["time"].dt.to_period("M")):
        monthly.append({"source": "lims", "point": point, "parameter": parameter, "month":str(month), "rows":len(group), "valid":int(group["value"].notna().sum()), "median":group["value"].median(), "above_10":int(group["value"].gt(10).sum()) if parameter == "Mg.Sulfur" else None})
pd.DataFrame(summary["lims"]).to_csv(OUT / "lims_profile.csv",index=False,encoding="utf-8-sig")
pd.DataFrame(monthly).to_csv(OUT / "monthly_profile.csv",index=False,encoding="utf-8-sig")

# Diagnostic agreement only: matching sample-time measurements is not an online feature join.
lab_sulfur = next(f for s,f in labs if s["parameter"]=="Mg.Sulfur").dropna(subset=["value"]).sort_values("time")
matched = pd.merge_asof(lab_sulfur[["time","value"]], pak["sulfur"][["time","value"]].dropna().sort_values("time"), on="time",direction="backward",tolerance=pd.Timedelta("10min"),suffixes=("_lab","_pak")).dropna()
errors = matched.value_pak - matched.value_lab
summary["sulfur_agreement_sample_time"] = {"n":len(matched),"mae":errors.abs().mean(),"bias_pak_minus_lab":errors.mean(),"correlation":matched.value_lab.corr(matched.value_pak), "opposite_sides_of_10":int((matched.value_lab.gt(10)!=matched.value_pak.gt(10)).sum())}
matched.to_csv(OUT / "sulfur_sample_time_comparison.csv",index=False,encoding="utf-8-sig")

ht = telemetry["242000"]
aligned = ht.join(pak["sulfur"].set_index("time")["value"].rename("pak_sulfur"),how="inner")
correlations = aligned.corr(numeric_only=True)["pak_sulfur"].drop("pak_sulfur").sort_values(key=abs,ascending=False)
summary["pak_sulfur_telemetry_correlations_diagnostic_only"] = correlations.to_dict()
dump("profile_summary.json", summary)
print("LIMS series",len(labs),"numeric measurements",sum(s["count"] for s,_ in labs))
print("LIMS sulfur",next(s for s,_ in labs if s["parameter"]=="Mg.Sulfur"))
print("SULFUR AGREEMENT",summary["sulfur_agreement_sample_time"])
print("SULFUR CORRELATIONS",correlations.head(8).to_dict())
