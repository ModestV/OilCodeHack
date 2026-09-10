"""Targeted evidence checks for anomalies and temporal source agreement."""
from pathlib import Path
import json
import sys
import argparse
import pandas as pd
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--source", required=True, type=Path)
parser.add_argument("--output", required=True, type=Path)
args = parser.parse_args()
SOURCE = args.source.resolve()
OUT = args.output.resolve()
OUT.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding="utf-8")
report = {}

def runs(frame, time, value, n=5):
    groups = frame[value].ne(frame[value].shift()).cumsum()
    result = frame.groupby(groups).agg(start=(time,"min"),end=(time,"max"),value=(value,"first"),rows=(value,"size")).sort_values("rows",ascending=False).head(n)
    result["elapsed_days"] = (result["end"]-result["start"]).dt.total_seconds()/86400
    return result.to_dict("records")

raw = pd.read_excel(next(SOURCE.glob("Выгрузка*")),header=None)
pak = pd.DataFrame({"time":pd.to_datetime(raw.iloc[2:,0],errors="coerce",format="mixed"),"value":pd.to_numeric(raw.iloc[2:,1],errors="coerce")}).dropna()
report["pak_sulfur_constant_runs"] = runs(pak,"time","value")
report["pak_sulfur_years"] = pak.assign(year=pak.time.dt.year,above=pak.value.gt(10)).groupby("year").agg(n=("value","size"),median=("value","median"),above_10=("above","sum")).reset_index().to_dict("records")

labraw = pd.read_excel(next(SOURCE.glob("ЛИМС*")),header=None)
report["lims_unparseable_rows"] = []
for col in range(0,len(labraw.columns),2):
    data = labraw.iloc[4:,[col,col+1]].copy()
    data.columns = ["date_raw","value_raw"]
    dates = pd.to_datetime(data.date_raw,errors="coerce",format="mixed")
    nums = pd.to_numeric(data.value_raw,errors="coerce")
    bad = data[(data.date_raw.notna() | data.value_raw.notna()) & (dates.isna() | nums.isna())]
    for index,row in bad.iterrows():
        report["lims_unparseable_rows"].append({"excel_cells": f"{get_column_letter(col+1)}{index+1}:{get_column_letter(col+2)}{index+1}", **row.to_dict()})

comparison = pd.read_csv(OUT / "sulfur_sample_time_comparison.csv",parse_dates=["time"])
report["lab_sulfur_largest"] = comparison.nlargest(8,"value_lab").to_dict("records")
report["sulfur_agreement_sensitivity"] = []
for cap in [None,100,20]:
    data = comparison if cap is None else comparison[comparison.value_lab.le(cap)]
    err = data.value_pak-data.value_lab
    report["sulfur_agreement_sensitivity"].append({"lab_upper_bound_for_diagnostic_only":cap,"n":len(data),"mae":err.abs().mean(),"median_absolute_error":err.abs().median(),"bias":err.mean(),"correlation":data.value_lab.corr(data.value_pak),"disagree_10":int((data.value_lab.gt(10)!=data.value_pak.gt(10)).sum())})
report["lab_sulfur_years"] = comparison.assign(year=comparison.time.dt.year,above=comparison.value_lab.gt(10)).groupby("year").agg(n=("value_lab","size"),median=("value_lab","median"),above_10=("above","sum")).reset_index().to_dict("records")
report["largest_sulfur_lab_gaps"] = comparison.assign(gap_hours=comparison.time.diff().dt.total_seconds()/3600,previous_time=comparison.time.shift()).nlargest(5,"gap_hours")[["previous_time","time","gap_hours"]].to_dict("records")

avt = pd.read_csv(SOURCE / "data" / "avt_tags.csv",parse_dates=["date"])
ht = pd.read_csv(SOURCE / "data" / "242000_tags.csv",parse_dates=["date"])
report["avt_D10_values"] = avt.D10.value_counts().to_dict()
report["avt_D10_constant_runs"] = runs(avt,"date","D10",3)
report["avt_F26_negative_count"] = int(avt.F26.lt(0).sum())
report["hydrotreating_selected_ranges"] = []
for tag in ["T5","T6","W7","P8","T11","T12","P13","F14","F15","T16","F17","T18","F19","Q20","Q21","F22","T23","P24","F25","F26"]:
    s = ht[tag]
    report["hydrotreating_selected_ranges"].append({"tag":tag,"min":s.min(),"median":s.median(),"max":s.max()})
report["telemetry_near_duplicate_pairs"] = []
for stage,frame in [("avt",avt),("242000",ht)]:
    signals = frame.drop(columns=[c for c in frame if c.startswith("Unnamed") or c=="date"])
    corr = signals.corr()
    for i in range(len(corr)):
        for j in range(i):
            if abs(corr.iloc[i,j]) > .995:
                report["telemetry_near_duplicate_pairs"].append({"stage":stage,"a":corr.index[i],"b":corr.index[j],"pearson":corr.iloc[i,j]})

vak_ht_t90 = 162.998 + .12945*ht.T12 + 59.57*ht.F15 + .00036*ht.W7 + .26366*ht.T23 - 424.72638*ht.F1/ht.F26.replace(0,float("nan"))
vak_avt_t50 = 981.06539 + .27467*avt.T42 - .32983*avt.F31/avt.F57.replace(0,float("nan")) - .49014*avt.T48
report["vak_literal_formula_diagnostics"] = [
    {"formula_cell":"ВАК!F2", "tag":"24-2000:GODT:T90", "median_literal_output":vak_ht_t90.median(), "note":"Literal coefficients and current CSV column names; not a validated model."},
    {"formula_cell":"ВАК!D2", "tag":"AVT6:350:T50", "median_literal_output":vak_avt_t50.median(), "note":"Decimal comma and x multiplication normalized literally; mapping not validated."},
]

profile = json.loads((OUT / "profile_summary.json").read_text(encoding="utf-8"))
report["counts"] = {"lab_series":len(profile["lims"]),"lab_numeric":sum(x["count"] for x in profile["lims"]),"lab_rows_with_time":sum(x["rows"] for x in profile["lims"]),"lab_declared":sum(x["declared_count"] for x in profile["lims"]),"telemetry_numeric_cells":len(avt)*97,"pak_above_10_fraction":pak.value.gt(10).mean(),"lab_above_10_fraction":comparison.value_lab.gt(10).mean()}
(OUT / "findings_checks.json").write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
print(json.dumps(report,ensure_ascii=False,indent=2,default=str))
