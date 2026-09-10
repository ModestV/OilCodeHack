"""Read source documents and inventory tabular files without modifying inputs."""
from pathlib import Path
import csv
import hashlib
import json
import sys
import argparse
from datetime import datetime
import openpyxl
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--source", required=True, type=Path)
parser.add_argument("--output", required=True, type=Path)
args = parser.parse_args()
SOURCE = args.source.resolve()
OUT = args.output.resolve()
OUT.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding="utf-8")

def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

inventory = []
for path in sorted(SOURCE.rglob("*")):
    if not path.is_file() or path.name.startswith("~$"):
        continue
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    inventory.append({"path": str(path.relative_to(SOURCE)), "bytes": path.stat().st_size, "sha256": digest.hexdigest()})
save("source_inventory.json", inventory)

for path in SOURCE.glob("*.docx"):
    if path.name.startswith("~$"):
        continue
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, table in enumerate(doc.tables, 1):
        lines.append(f"\nTABLE {i}")
        lines.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
    (OUT / "requirements_extracted.txt").write_text("\n".join(lines), encoding="utf-8")
    print("DOCX", path.name, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))

for path in SOURCE.glob("*.pdf"):
    pdf = PdfReader(path)
    lines = []
    for i, page in enumerate(pdf.pages, 1):
        lines.append(f"\nPAGE {i}\n{page.extract_text()}")
        print("PDF page", i, "size", page.mediabox, "images", len(page.images))
    (OUT / "schemes_extracted.txt").write_text("\n".join(lines), encoding="utf-8")

books = []
for path in SOURCE.glob("*.xlsx"):
    if path.name.startswith("~$"):
        continue
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    info = {"file": path.name, "sheets": []}
    for sheet in book:
        rows = []
        nonempty = 0
        tail = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if not any(x is not None for x in row):
                continue
            nonempty += 1
            record = {"excel_row": index, "values": row}
            if len(rows) < 12:
                rows.append(record)
            tail = (tail + [record])[-3:]
        info["sheets"].append({"sheet": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "nonempty_rows": nonempty, "head": rows, "tail": tail})
    books.append(info)
    book.close()
save("workbooks_structure.json", books)

samples = []
for path in (SOURCE / "data").glob("*.csv"):
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(20000)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        stream.seek(0)
        reader = csv.reader(stream, dialect)
        rows = [next(reader) for _ in range(6)]
    samples.append({"file": path.name, "delimiter": dialect.delimiter, "head": rows})
save("csv_structure.json", samples)
print("Saved source inventory, extracted text, workbook structure and CSV samples to", OUT)
