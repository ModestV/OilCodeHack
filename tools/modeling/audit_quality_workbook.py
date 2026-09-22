"""Read-only original XLSX audit; run with the bundled Python/openpyxl environment."""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def run(path, output):
    sheet=load_workbook(path,read_only=True,data_only=True).worksheets[0]
    rows=list(sheet.iter_rows(values_only=True)); result={}
    for name,col in [('t95',92),('cetane',102)]:
        values=[{'row':i+1,'time':str(row[col]),'value':row[col+1]}
                for i,row in enumerate(rows) if i>=4 and row[col] is not None and isinstance(row[col+1],(float,int))]
        seen={}; duplicates=[]
        for v in values:
            if v['time'] in seen:duplicates.append([seen[v['time']],v])
            seen[v['time']]=v
        result[name]=dict(header_count=rows[3][col+1],numeric_rows=len(values),unique_timestamps=len(seen),
                          duplicate_rows=duplicates,headers=[[str(x) for x in row[col:col+2]] for row in rows[:4]])
    result.update(source_file=str(path),sheet=sheet.title,sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source',type=Path,required=True)
    p.add_argument('--output',type=Path,default=Path('reports/modeling/cetane-t95/source-workbook-audit.json'))
    a=p.parse_args();run(a.source,a.output)
